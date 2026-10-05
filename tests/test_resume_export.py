import os

import pytest
from docx import Document

from app.db.connection import get_db
from app.db.models import JobStatus, NormalizedJob
from app.db.repository import JobRepository
from app.profile.models import (
    FactBank,
    FactCategory,
    FactItem,
)
from app.resume.engine import ResumeTailoringEngine
from app.resume.export.docx import DOCXResumeExporter
from app.resume.export.models import ExportFormat
from app.resume.export.pdf import PDFResumeExporter
from app.resume.export.validation import ResumeExportValidator


@pytest.fixture
def db_conn():
    with get_db(":memory:") as conn:
        yield conn


@pytest.fixture
def sample_fact_bank():
    return FactBank(
        facts=[
            FactItem(
                fact_id="EDU-001",
                category=FactCategory.EDUCATION,
                subject="B.Tech in Electronics and Communication Engineering",
                value={"degree": "Bachelor of Technology", "institution": "National Institute of Technology", "graduation_year": 2025, "gpa": "7.38/10"},
                verified=True,
            ),
            FactItem(
                fact_id="SKILL-001",
                category=FactCategory.SKILL,
                subject="SystemVerilog",
                value={"skill_name": "SystemVerilog", "proficiency": "Advanced"},
                verified=True,
            ),
            FactItem(
                fact_id="SKILL-002",
                category=FactCategory.SKILL,
                subject="UVM",
                value={"skill_name": "UVM", "proficiency": "Intermediate"},
                verified=True,
            ),
            FactItem(
                fact_id="SKILL-003",
                category=FactCategory.SKILL,
                subject="AXI Protocol",
                value={"skill_name": "AXI4 Protocol", "proficiency": "Intermediate"},
                verified=True,
            ),
            FactItem(
                fact_id="PROJ-001",
                category=FactCategory.PROJECT,
                subject="AXI4 Interface Verification IP",
                value={
                    "title": "AXI4 UVC Testbench Environment",
                    "technologies": ["SystemVerilog", "UVM", "QuestaSim", "SVA"],
                    "role": "Lead Verification Developer",
                    "bullets": [
                        "Architected modular UVM testbench comprising active master agent, monitor, and scoreboard.",
                        "Developed constrained-random sequence library generating single, incrementing, and wrap burst transfers.",
                        "Implemented SystemVerilog concurrent assertions (SVA) verifying VALID/READY handshakes.",
                    ],
                },
                verified=True,
            ),
        ]
    )


@pytest.fixture
def target_job(db_conn):
    job_repo = JobRepository(db_conn)
    job = NormalizedJob(
        company="Texas Instruments",
        title="Design Verification Engineer",
        location="Bengaluru",
        country="India",
        skills=["SystemVerilog", "UVM", "AXI Protocol", "SVA"],
        source="manual",
        status=JobStatus.ACTIVE,
        first_seen="2026-10-04T10:00:00Z",
        last_seen="2026-10-04T10:00:00Z",
        fingerprint="ti-dv-export-test-01",
    )
    job_id = job_repo.insert_normalized_job(job)
    job.id = job_id
    return job


def test_docx_resume_export_and_validation(tmp_path, db_conn, sample_fact_bank, target_job):
    """Test generating, exporting, and validating an ATS-safe DOCX file."""
    engine = ResumeTailoringEngine(db_conn, fact_bank=sample_fact_bank)
    resume = engine.generate_tailored_resume(job_id=target_job.id)

    docx_exporter = DOCXResumeExporter(output_dir=str(tmp_path))
    docx_path = docx_exporter.export(resume, filename="test_resume.docx")

    assert os.path.exists(docx_path)

    # Inspect docx structure directly
    doc = Document(docx_path)
    assert len(doc.tables) == 0  # ATS invariant: No tables
    assert len(doc.paragraphs) > 5

    # Run validator
    validator = ResumeExportValidator()
    report = validator.validate_export(docx_path, ExportFormat.DOCX, resume)

    assert report.status == "PASS"
    assert report.is_export_validated is True
    assert len(report.missing_sections) == 0
    assert len(report.content_mismatches) == 0
    assert report.extracted_text_length > 200


def test_pdf_resume_export_and_validation(tmp_path, db_conn, sample_fact_bank, target_job):
    """Test generating, exporting, and validating an ATS-safe selectable PDF file."""
    engine = ResumeTailoringEngine(db_conn, fact_bank=sample_fact_bank)
    resume = engine.generate_tailored_resume(job_id=target_job.id)

    pdf_exporter = PDFResumeExporter(output_dir=str(tmp_path))
    pdf_path = pdf_exporter.export(resume, filename="test_resume.pdf")

    assert os.path.exists(pdf_path)

    # Run validator
    validator = ResumeExportValidator()
    report = validator.validate_export(pdf_path, ExportFormat.PDF, resume)

    assert report.status == "PASS"
    assert report.is_export_validated is True
    assert len(report.missing_sections) == 0
    assert len(report.content_mismatches) == 0
    assert report.extracted_text_length > 200


def test_export_validator_rejects_missing_sections(tmp_path, sample_fact_bank, target_job, db_conn):
    """Test validator catches corrupted exports missing mandatory ATS sections."""
    engine = ResumeTailoringEngine(db_conn, fact_bank=sample_fact_bank)
    resume = engine.generate_tailored_resume(job_id=target_job.id)

    # Write a dummy txt file lacking sections
    corrupt_txt = tmp_path / "corrupt.txt"
    corrupt_txt.write_text("Candidate Name\nJust a simple text file without proper sections.\n" * 10, encoding="utf-8")

    validator = ResumeExportValidator()
    report = validator.validate_export(str(corrupt_txt), ExportFormat.PLAINTEXT, resume)

    assert report.status == "FAIL"
    assert report.is_export_validated is False
    assert len(report.missing_sections) >= 3


def test_export_validator_detects_introduced_fabricated_metrics(tmp_path, sample_fact_bank, target_job, db_conn):
    """Test validator catches fabricated metric patterns injected into exported files."""
    engine = ResumeTailoringEngine(db_conn, fact_bank=sample_fact_bank)
    resume = engine.generate_tailored_resume(job_id=target_job.id)

    # Inject a fabricated metric into exported text
    fabricated_txt = tmp_path / "fabricated.txt"
    text_content = (
        "CANDIDATE NAME\n\n"
        "PROFESSIONAL SUMMARY\nExperienced in verification.\n\n"
        "TECHNICAL SKILLS\nSkills: SystemVerilog, UVM\n\n"
        "TECHNICAL PROJECTS\nAXI4 UVC Testbench\n* Achieved 99% coverage and fixed 45 bugs.\n\n"
        "EDUCATION\nB.Tech - National Institute of Technology (Graduation: 2025)\n"
    )
    fabricated_txt.write_text(text_content, encoding="utf-8")

    validator = ResumeExportValidator()
    report = validator.validate_export(str(fabricated_txt), ExportFormat.PLAINTEXT, resume)

    assert report.status == "FAIL"
    assert report.is_export_validated is False
    assert any("Fabricated metric" in err for err in report.errors)
